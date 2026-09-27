import os
import sys
import sqlite3
import subprocess

# 1. Ensure pandas and openpyxl are installed
try:
    import pandas as pd
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing spreadsheet dependencies (pandas, openpyxl)...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "pandas", "openpyxl"])
    import pandas as pd
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

def export_sqlite_to_xlsx():
    db_path = os.path.join(".agents", "governance", "governance.db")
    xlsx_path = os.path.join("artifacts", "primecare_governance_registry.xlsx")
    
    if not os.path.exists(db_path):
        print(f"Error: SQLite database not found at {db_path}")
        sys.exit(1)
        
    os.makedirs(os.path.dirname(xlsx_path), exist_ok=True)
    
    # Connect to SQLite
    print(f"Reading database: {db_path}...")
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # Fetch all table names
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
    tables = [row[0] for row in cursor.fetchall()]
    
    print(f"Identified tables to export: {tables}")
    
    # Initialize Excel Writer
    with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
        for table in tables:
            print(f"Exporting table '{table}'...")
            df = pd.read_sql_query(f"SELECT * FROM {table}", conn)
            
            # Write to excel sheet
            df.to_excel(writer, sheet_name=table, index=False)
            
            # Access openpyxl sheet to apply high-fidelity styling
            workbook = writer.book
            worksheet = writer.sheets[table]
            
            # --- Premium Styling System ---
            # Font Styles
            header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            data_font = Font(name='Segoe UI', size=10, color='1E293B')
            
            # Fills
            header_fill = PatternFill(start_color='0284C7', end_color='0284C7', fill_type='solid') # Sky Blue theme
            zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid') # Slate 50
            
            # Alignments
            left_align = Alignment(horizontal='left', vertical='center')
            center_align = Alignment(horizontal='center', vertical='center')
            right_align = Alignment(horizontal='right', vertical='center')
            
            # Borders
            thin_side = Side(border_style="thin", color="E2E8F0")
            cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            
            # Format Headers
            worksheet.row_dimensions[1].height = 28
            for col_idx, col_name in enumerate(df.columns, 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = cell_border
                
            # Format Data Rows (optimized for extremely large tables)
            if len(df) < 3000:
                for row_idx in range(2, worksheet.max_row + 1):
                    worksheet.row_dimensions[row_idx].height = 20
                    is_even = (row_idx % 2 == 0)
                    
                    for col_idx in range(1, worksheet.max_column + 1):
                        cell = worksheet.cell(row=row_idx, column=col_idx)
                        cell.font = data_font
                        cell.border = cell_border
                        
                        if is_even:
                            cell.fill = zebra_fill
                            
                        # Deduce alignment based on data type
                        val = cell.value
                        if isinstance(val, (int, float)):
                            cell.alignment = right_align
                        elif str(val).lower() in ('true', 'false'):
                            cell.alignment = center_align
                        else:
                            cell.alignment = left_align
            else:
                print(f"  Skipping cell-by-cell styling for large table '{table}' ({len(df)} rows) to speed up compilation...")
            
            # Auto-fit columns (optimized for extremely large tables)
            if len(df) < 3000:
                for col in worksheet.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        val_str = str(cell.value or '')
                        if len(val_str) > max_len:
                            max_len = len(val_str)
                    # Pad for readability and filter out huge values
                    worksheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)
            else:
                for col_idx, col_name in enumerate(df.columns, 1):
                    col_letter = get_column_letter(col_idx)
                    worksheet.column_dimensions[col_letter].width = 20
                
    conn.close()
    print(f"[SUCCESS] Generated beautiful stylized workbook at: {xlsx_path}")
    
    # Duplicate to .agents/governance/primecare_governance_audit.xlsx
    import shutil
    audit_path = os.path.join(".agents", "governance", "primecare_governance_audit.xlsx")
    os.makedirs(os.path.dirname(audit_path), exist_ok=True)
    shutil.copyfile(xlsx_path, audit_path)
    print(f"[SUCCESS] Duplicated stylized workbook to: {audit_path}")

if __name__ == "__main__":
    export_sqlite_to_xlsx()

